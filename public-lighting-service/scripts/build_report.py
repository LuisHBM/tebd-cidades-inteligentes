"""Gera o relatório no template SBC com os dados e diagramas do projeto."""
import json,pathlib,re
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
PROJECT_DIR=pathlib.Path(__file__).resolve().parents[1]
manifest=json.loads((PROJECT_DIR/'docs/data_manifest.json').read_text())
validation=json.loads((PROJECT_DIR/'docs/validation/postgresql_validation.json').read_text())
assert manifest['counts']['solicitacoes']==validation['counts']['olap.fato_solicitacao']==34492
from build_diagrams import build
for kind in ['oltp','olap']:
 build(kind)
document=Document(PROJECT_DIR/'references/sbc_template.docx')
for e in list(document._element.body):
 if e.tag!=qn('w:sectPr'):document._element.body.remove(e)
document.core_properties.title='Atendimento de iluminação pública: banco de dados para análise das solicitações do Recife'
document.core_properties.author='';document.core_properties.last_modified_by='';document.core_properties.comments=''
normal=document.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal.paragraph_format.space_before=Pt(6);normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.line_spacing=1
for n in ['Heading 1','Heading 2']:
 st=document.styles[n];st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
 st.paragraph_format.space_before=Pt(12)
 st.paragraph_format.space_after=Pt(12)
 st.paragraph_format.keep_with_next=True
 if st._element.pPr is not None:
  num=st._element.pPr.find(qn('w:numPr'))
  if num is not None:st._element.pPr.remove(num)
CODE_IDENTIFIERS={'grupo_servico','solicitacao','solicitacao_id','quantidade'}
IDENTIFIER_PATTERN=re.compile(r'\b('+'|'.join(sorted(CODE_IDENTIFIERS,key=len,reverse=True))+r')\b')
def add_formatted_text(paragraph,text):
 for part in IDENTIFIER_PATTERN.split(text):
  if part:
   run=paragraph.add_run(part)
   if part in CODE_IDENTIFIERS:run.italic=True

def add_paragraph(text,style=None,indent=False):
 p=document.add_paragraph(style=style)
 add_formatted_text(p,text)
 p.paragraph_format.first_line_indent=Cm(1.27 if indent else 0)
 p.paragraph_format.keep_together=True
 if not style:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 return p
def add_heading(t):return add_paragraph(t,'Heading 1')
def add_figure(kind,caption):
 p=add_paragraph('', 'Figure');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(PROJECT_DIR/'docs/diagrams'/f'{kind}_model.png'),width=Cm(15))
 p=add_paragraph(caption,'Legenda');p.paragraph_format.keep_together=True
 p.paragraph_format.space_after=Pt(6)
for text,style in [
 ('Atendimento de iluminação pública: análise das solicitações do Recife','Heading'),
 ('Luís Henrique, Alan Vitor e Kaique William','Author'),('Universidade do Estado da Bahia (UNEB)','Address')]:add_paragraph(text,style)

add_heading('1. Introdução')
add_paragraph('Os registros de atendimento da iluminação pública do Recife são publicados pela Central 156 em um arquivo CSV. A análise desse arquivo exige organizar e relacionar os dados para identificar a concentração de solicitações por bairro e serviço e comparar a situação publicada e o intervalo até o atendimento registrado.')
add_paragraph('A solução proposta é um banco de dados relacional com carga reproduzível e um modelo multidimensional. Essa estrutura permite consultar os registros por território, serviço e período para apoiar a seleção de demandas que precisam de revisão pela gestão municipal.',indent=True)
add_paragraph('A proposta se relaciona às metas 11.7 e 11.3 do Objetivo de Desenvolvimento Sustentável 11 (ODS 11), que tratam do acesso a espaços públicos e da gestão urbana. As metas são apresentadas pelo Instituto de Pesquisa Econômica Aplicada (Ipea) em seu portal sobre os ODS [Ipea, 2026]. A iluminação faz parte das condições de uso de ruas e espaços públicos. As informações de atendimento permitem estudar esse serviço, mas não medem os indicadores oficiais das metas nem demonstram efeitos sobre segurança, consumo de energia ou emissões.',indent=True)
add_heading('2. Dados e requisitos')
add_paragraph('Os dados vêm do CSV da Central 156 da Autarquia de Manutenção e Limpeza Urbana do Recife (EMLURB) [Emlurb, 2026]. A base está disponível no portal de dados abertos municipal sob licença Open Database License (ODbL) [Recife, 2026]. O recorte reúne 34.492 registros do grupo ILUMINACAO PUBLICA, com demandas abertas de 01/01 a 31/08/2025. O arquivo foi baixado diretamente em 18/10/2026, sem scraping. Um manifesto documenta a origem, as transformações e o hash SHA-256 que identifica a versão utilizada.')
add_paragraph('Os requisitos funcionais (RF) do sistema são:').paragraph_format.keep_with_next=True
requirements=[
 'Importar o recorte de iluminação pública e registrar a origem dos dados.',
 'Relacionar região político-administrativa (RPA), bairro, grupo de serviço, serviço, situação e solicitação.',
 'Preservar o endereço e as datas de abertura e de última situação.',
 'Rejeitar chaves duplicadas, referências inexistentes e datas invertidas.',
 'Consultar volume e situação por bairro, RPA, serviço e mês de abertura.',
 'Calcular o intervalo até a última situação, separando ATENDIDA das demais.',
 'Refazer a carga analítica com o mesmo conteúdo a partir da fonte selecionada.'
]
table=document.add_table(rows=1,cols=2)
table.alignment=WD_TABLE_ALIGNMENT.LEFT
table.autofit=False
widths=[Cm(1.5),Cm(13.5)]
for column,width in zip(table.columns,widths):column.width=width
for cell,text in zip(table.rows[0].cells,['Código','Descrição']):cell.text=text
for i,text in enumerate(requirements,1):
 cells=table.add_row().cells
 cells[0].text=f'RF{i:02d}'
 cells[1].text=text
props=table._tbl.tblPr
width=props.find(qn('w:tblW'));width.set(qn('w:type'),'dxa');width.set(qn('w:w'),str(Cm(15).twips))
borders=OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
 edge=OxmlElement('w:'+side)
 for key,value in [('val','single'),('sz','4'),('color','000000')]:edge.set(qn('w:'+key),value)
 borders.append(edge)
props.append(borders)
margins=OxmlElement('w:tblCellMar')
for side,value in [('top','10'),('bottom','10'),('left','40'),('right','40')]:
 margin=OxmlElement('w:'+side);margin.set(qn('w:w'),value);margin.set(qn('w:type'),'dxa');margins.append(margin)
props.append(margins)
for ri,row in enumerate(table.rows):
 trprops=row._tr.get_or_add_trPr();trprops.append(OxmlElement('w:cantSplit'))
 if ri==0:trprops.append(OxmlElement('w:tblHeader'))
 for ci,cell in enumerate(row.cells):
  cell.width=widths[ci]
  cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for paragraph in cell.paragraphs:
   text=paragraph.text
   paragraph.clear()
   add_formatted_text(paragraph,text)
   paragraph.style=document.styles['Normal']
   paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
   pf=paragraph.paragraph_format
   pf.left_indent=Cm(0);pf.right_indent=Cm(0);pf.first_line_indent=Cm(0)
   pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1
   pf.keep_together=True
   pf.keep_with_next=ri<len(table.rows)-1
   for run in paragraph.runs:
    run.font.name='Times New Roman';run.font.size=Pt(12);run.bold=ri==0
add_paragraph('Cadastro de cidadãos, equipes, ordens de serviço e inventário de postes ficam fora do escopo.',indent=True)
add_heading('3. Modelagem dos dados')
add_paragraph('O modelo transacional (OLTP, Online Transaction Processing) tem seis entidades (Figura 1). A RPA reúne bairros, enquanto grupo_servico reúne os tipos de serviço. A entidade solicitacao associa bairro, serviço e situação ao endereço e às datas da demanda. O cadastro de situações contém os estados publicados pela fonte. Como o recorte abrange apenas iluminação, grupo_servico tem um único registro. Todos esses domínios existem no CSV. As entidades mostram as chaves primárias (PK) e estrangeiras (FK).')
add_figure('oltp','Figura 1. Diagrama entidade-relacionamento (DER) do modelo OLTP.')
add_paragraph('O CSV não contém protocolo público. Por isso, solicitacao_id usa a posição da linha de dados no arquivo original, identificado pelo hash. A chave vale apenas nessa versão. O banco armazena a última situação disponível, sem criar etapas de atendimento que a fonte não informa. As restrições SQL rejeitam datas invertidas, chaves duplicadas e referências a cadastros inexistentes.',indent=True)
add_paragraph('O modelo analítico (OLAP, Online Analytical Processing) segue um esquema estrela, com um registro na tabela fato para cada solicitação publicada (Figura 2). As quatro dimensões são data, bairro, serviço e situação. A dimensão de data é consultada em dois papéis: abertura da demanda e última situação. RPA é um atributo de bairro, e grupo é um atributo de serviço.',indent=True)
add_figure('olap','Figura 2. DER do modelo OLAP com duas relações para a dimensão de data.')
add_paragraph('A tabela fato contém a quantidade, fixada em 1 por registro, e o intervalo em dias entre a demanda e a última situação. A análise do intervalo até o atendimento deve considerar somente os registros ATENDIDA. Para as demais situações, esse valor indica apenas o tempo até a movimentação publicada. Mesmo em ATENDIDA, ele não comprova quanto durou o reparo físico.',indent=True)
add_heading('4. Implementação e validação')
add_paragraph('A implementação usa PostgreSQL 16 e Python 3. O tratamento retira espaços das extremidades dos textos, mantém os códigos oficiais de RPA, grupo e serviço e cria chaves locais para bairros e situações. Valores ausentes permanecem sem preenchimento, e registros no mesmo endereço são mantidos. A análise territorial usa bairro e RPA. As coordenadas ficaram fora do modelo porque faltam em 31.127 registros.')
add_paragraph('O roteiro SQL criou e carregou seis tabelas OLTP, quatro dimensões e uma tabela fato. A carga contém 34.492 registros, distribuídos por 93 bairros, seis RPAs, dez serviços e cinco situações. A comparação entre os modelos, feita por identificador, chaves, datas e medida, não encontrou divergências. Os testes confirmaram a rejeição de PK duplicada, FK inexistente e última situação anterior à demanda. Repetir a carga OLAP preservou o conteúdo.',indent=True)
add_paragraph('Na versão analisada, 33.662 registros estão como ATENDIDA e 830 têm outras situações. Entre os registros ATENDIDA, a mediana do intervalo até esse status é de dois dias. Há datas de última situação até 10/09/2025, posteriores ao fim do período de abertura selecionado. Assim, os 830 registros descrevem o estado publicado no arquivo, não o estoque pendente em 31/08/2025 ou o atendimento atual do município.',indent=True)
add_paragraph('O dicionário de dados documenta as 42 colunas dos modelos, com tipos, obrigatoriedade de preenchimento, chaves e restrições. Os scripts SQL permitem criar os dois ambientes e reproduzir a carga a partir do recorte selecionado. As conferências de integridade e a identificação da fonte permitem verificar os resultados apresentados.',indent=True)
add_heading('5. Limites e continuidade')
add_paragraph('As consultas ajudam a selecionar bairros e serviços para investigação. A base, porém, não informa o tamanho da rede, os custos ou os prazos contratuais de atendimento. Também não identifica cada poste nem permite saber se dois registros se referem ao mesmo defeito. Por isso, o projeto não estima taxas de falha por ponto de luz. O OLTP é um protótipo construído com dados públicos, sem integração com o sistema interno da prefeitura.')
add_paragraph('Para a entrega de 06/11/2026, o trabalho prevê dashboards operacional, tático e estratégico, com cinco consultas em cada nível, um agrupamento exploratório de bairros e um plano de ação 5W2H para definir ações, justificativas, locais, prazos, responsáveis, procedimentos e custos. Os perfis encontrados servirão de base para propor a revisão de registros e a seleção de locais para verificação. A aplicação dessas ações dependerá de confirmação com o órgão responsável.',indent=True)
add_heading('Referências')
for text in [
 'Emlurb (2026). Sobre. Autarquia de Manutenção e Limpeza Urbana do Recife. https://emlurb.recife.pe.gov.br/sobre-0. Acesso em 18 out. 2026.',
 'Instituto de Pesquisa Econômica Aplicada (Ipea) (2026). ODS 11: Cidades e comunidades sustentáveis. https://www.ipea.gov.br/ods/ods11.html. Acesso em 18 out. 2026.',
 'Recife (2026). Central de Atendimento de Serviços da EMLURB 156. Recurso 2025, Prefeitura do Recife. https://dados.recife.pe.gov.br/dataset/central-de-atendimento-de-servicos-da-emlurb-156. Acesso em 18 out. 2026.'
]:
 p=add_paragraph(text);p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
document.save(PROJECT_DIR/'reports/relatorio_parcial_2026_10_19.docx')
print(PROJECT_DIR/'reports/relatorio_parcial_2026_10_19.docx')
